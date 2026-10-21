"""Carga de datos de demostración desde la planilla de ejemplo.

Permite correr MerchBot sin Google Sheets: lee `data/Stock_Merchandising_Demo.xlsx`
(la misma plantilla que se sube a Sheets) y llena la base local.

La planilla calcula "Pedidos", "Stock Actual" y "Estado" con fórmulas. openpyxl
no evalúa fórmulas, así que esos valores se recalculan acá con la misma lógica:
    Pedidos      = suma de cantidades de la hoja Pedidos para ese SKU
    Stock Actual = Stock Inicial + Ingresos - Pedidos
    Estado       = "OK" si Stock Actual > 0, si no "COMPRAR"
"""
from __future__ import annotations

import logging
from collections import defaultdict
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from merchbot.config import settings
from merchbot.data.sheets_sync import (
    SyncResult,
    _parse_fecha,
    _replace_pedidos,
    _safe_int,
    _upsert_stock,
)
from merchbot.timeutils import utc_now


logger = logging.getLogger(__name__)


def _filas(hoja) -> list[dict]:
    """Filas de una hoja como dicts {encabezado: valor}, salteando las vacías."""
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return []
    encabezados = [str(h).strip() if h is not None else "" for h in filas[0]]
    resultado = []
    for fila in filas[1:]:
        if all(v is None or str(v).strip() == "" for v in fila):
            continue
        resultado.append(dict(zip(encabezados, fila)))
    return resultado


def leer_planilla_demo(path: Path) -> tuple[list[dict], list[dict]]:
    """Devuelve (stock, pedidos) listos para cargar en la base."""
    libro = load_workbook(path, read_only=True, data_only=False)
    try:
        filas_pedidos = _filas(libro[settings.sheets_pedidos_tab])
        filas_stock = _filas(libro[settings.sheets_stock_tab])
    finally:
        libro.close()

    pedidos = []
    vendidos_por_sku: dict[str, int] = defaultdict(int)
    for fila in filas_pedidos:
        sku = str(fila.get("SKU") or "").strip()
        cantidad = _safe_int(fila.get("Cantidad"))
        fecha = fila.get("Fecha")
        fecha = fecha.date() if hasattr(fecha, "date") else _parse_fecha(fecha)
        if not sku or cantidad == 0:
            continue
        vendidos_por_sku[sku] += cantidad
        pedidos.append({
            "fecha": fecha,
            "cliente": str(fila.get("Cliente") or "").strip(),
            "sku": sku,
            "producto": str(fila.get("Producto") or "").strip(),
            "cantidad": cantidad,
        })

    stock = []
    for fila in filas_stock:
        sku = str(fila.get("SKU") or "").strip()
        if not sku:
            continue
        inicial = _safe_int(fila.get("Stock Inicial"))
        ingresos = _safe_int(fila.get("Ingresos"))
        vendidos = vendidos_por_sku.get(sku, 0)
        actual = inicial + ingresos - vendidos
        stock.append({
            "sku": sku,
            "producto": str(fila.get("Producto") or "").strip(),
            "modelo": str(fila.get("Modelo") or "").strip(),
            "color": str(fila.get("Color") or "").strip(),
            "talle": str(fila.get("Talle") or "").strip(),
            "stock_inicial": inicial,
            "ingresos": ingresos,
            "pedidos": vendidos,
            "stock_actual": actual,
            "estado": "OK" if actual > 0 else "COMPRAR",
        })
    return stock, pedidos


def cargar_datos_demo(path: Optional[Path] = None) -> SyncResult:
    """Carga la planilla de demo en la base. Es idempotente: se puede correr varias veces."""
    path = Path(path or settings.demo_xlsx_path)
    now = utc_now()
    if not path.exists():
        return SyncResult(ok=False, error=f"No encontré la planilla de demo en {path}", timestamp=now)
    try:
        stock, pedidos = leer_planilla_demo(path)
        insertados, actualizados = _upsert_stock(stock, now)
        cargados = _replace_pedidos(pedidos, now)
    except Exception as e:  # noqa: BLE001 — se reporta al usuario, no se cae la app
        logger.exception("No pude cargar los datos de demo")
        return SyncResult(ok=False, error=f"{type(e).__name__}: {e}", timestamp=now)

    logger.info("Datos de demo cargados: %d SKUs, %d pedidos", len(stock), cargados)
    return SyncResult(
        ok=True,
        stock_total=len(stock),
        stock_insertados=insertados,
        stock_actualizados=actualizados,
        pedidos_total=len(pedidos),
        pedidos_insertados=cargados,
        timestamp=now,
    )

