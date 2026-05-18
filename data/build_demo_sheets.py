"""Genera el XLSX de demo para el bot de merchandising.

Estructura (igual al Google Sheets real de la empresa):
- Hoja 'Stock': inventario con fórmulas
    Stock Actual = Stock Inicial + Ingresos - Pedidos
    + Estado condicional (OK/COMPRAR) + formato verde/rojo
- Hoja 'Pedidos': histórico de pedidos. Los pedidos suman al cálculo de stock via SUMIF.

Para regenerar:
    python data/build_demo_sheets.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule

# ---------------------------------------------------------------------------
# Productos base (37 SKUs realistas)
# Tupla: (sku, producto, modelo, color, talle, stock_inicial, ingresos)
# ---------------------------------------------------------------------------
PRODUCTOS = [
    # REMERAS cuello redondo
    ("REM-CR-BCO-S",  "Remera",   "Cuello Redondo", "Blanco", "S",   80, 50),
    ("REM-CR-BCO-M",  "Remera",   "Cuello Redondo", "Blanco", "M",  100, 100),
    ("REM-CR-BCO-L",  "Remera",   "Cuello Redondo", "Blanco", "L",   80, 80),
    ("REM-CR-BCO-XL", "Remera",   "Cuello Redondo", "Blanco", "XL",  50, 30),
    ("REM-CR-NEG-M",  "Remera",   "Cuello Redondo", "Negro",  "M",   60, 40),
    ("REM-CR-NEG-L",  "Remera",   "Cuello Redondo", "Negro",  "L",   40, 20),
    # REMERAS dry-fit
    ("REM-DF-NEG-M",  "Remera",   "Dry-Fit",        "Negro",  "M",   30, 0),
    ("REM-DF-NEG-L",  "Remera",   "Dry-Fit",        "Negro",  "L",   25, 0),
    ("REM-DF-AZU-M",  "Remera",   "Dry-Fit",        "Azul",   "M",   20, 30),
    ("REM-DF-AZU-L",  "Remera",   "Dry-Fit",        "Azul",   "L",   15, 20),
    ("REM-DF-ROJ-M",  "Remera",   "Dry-Fit",        "Rojo",   "M",   10, 0),
    # BUZOS canguro
    ("BUZ-CN-GRI-M",  "Buzo",     "Canguro",        "Gris",   "M",   40, 30),
    ("BUZ-CN-GRI-L",  "Buzo",     "Canguro",        "Gris",   "L",   50, 50),
    ("BUZ-CN-GRI-XL", "Buzo",     "Canguro",        "Gris",   "XL",  30, 20),
    ("BUZ-CN-NEG-M",  "Buzo",     "Canguro",        "Negro",  "M",   35, 25),
    ("BUZ-CN-NEG-L",  "Buzo",     "Canguro",        "Negro",  "L",   45, 30),
    ("BUZ-CN-NEG-XL", "Buzo",     "Canguro",        "Negro",  "XL",  20, 15),
    ("BUZ-CN-BOR-L",  "Buzo",     "Canguro",        "Bordó",  "L",   15, 0),
    # BUZOS cuello redondo
    ("BUZ-CR-GRI-M",  "Buzo",     "Cuello Redondo", "Gris",   "M",   25, 0),
    ("BUZ-CR-GRI-L",  "Buzo",     "Cuello Redondo", "Gris",   "L",   30, 20),
    ("BUZ-CR-NEG-M",  "Buzo",     "Cuello Redondo", "Negro",  "M",   20, 10),
    # CAMPERAS con cierre
    ("CAM-CIE-NEG-M", "Campera",  "Con Cierre",     "Negro",  "M",   18, 12),
    ("CAM-CIE-NEG-L", "Campera",  "Con Cierre",     "Negro",  "L",   20, 10),
    ("CAM-CIE-AZU-L", "Campera",  "Con Cierre",     "Azul",   "L",   12, 8),
    # CHOMBAS pique
    ("CHO-PIQ-BCO-M", "Chomba",   "Piqué",          "Blanco", "M",   25, 25),
    ("CHO-PIQ-BCO-L", "Chomba",   "Piqué",          "Blanco", "L",   30, 30),
    ("CHO-PIQ-AZM-L", "Chomba",   "Piqué",          "Azul Marino", "L", 20, 15),
    ("CHO-PIQ-NEG-M", "Chomba",   "Piqué",          "Negro",  "M",   18, 12),
    # GORRAS (sin talle real, único)
    ("GOR-RES-NEG",   "Gorra",    "Con Respiración","Negro",  "Único", 60, 40),
    ("GOR-RES-BCO",   "Gorra",    "Con Respiración","Blanco", "Único", 50, 30),
    ("GOR-RES-AZU",   "Gorra",    "Con Respiración","Azul",   "Único", 40, 20),
    ("GOR-RES-ROJ",   "Gorra",    "Con Respiración","Rojo",   "Único", 20, 0),
    # TAZAS (sin talle)
    ("TAZ-CER-BCA",   "Taza",     "Cerámica",       "Blanco", "Único", 200, 150),
    ("TAZ-CER-NEG",   "Taza",     "Cerámica Interior Negro", "Negro", "Único", 80, 50),
    ("TAZ-MAG",       "Taza",     "Mágica",         "Negro/Blanco", "Único", 50, 0),
    # LLAVEROS
    ("LLA-ACR-RED",   "Llavero",  "Acrílico Redondo","Transparente", "Único", 300, 200),
    ("LLA-ACR-REC",   "Llavero",  "Acrílico Rectangular","Transparente", "Único", 250, 150),
]

# ---------------------------------------------------------------------------
# Pedidos recientes (clientes ficticios típicos del sector)
# Tupla: (fecha, cliente, sku, cantidad)
# Los pedidos descuentan del stock via SUMIF en la hoja Stock
# ---------------------------------------------------------------------------
PEDIDOS = [
    ("2026-04-15", "Coca-Cola - Evento Verano",  "REM-CR-BCO-M",  120),
    ("2026-04-15", "Coca-Cola - Evento Verano",  "REM-CR-BCO-L",   80),
    ("2026-04-15", "Coca-Cola - Evento Verano",  "GOR-RES-ROJ",    40),
    ("2026-04-18", "Banco Nación - Capacitación","REM-CR-BCO-M",   60),
    ("2026-04-18", "Banco Nación - Capacitación","CHO-PIQ-BCO-L",  40),
    ("2026-04-20", "Universidad UADE",           "BUZ-CN-GRI-L",   70),
    ("2026-04-20", "Universidad UADE",           "BUZ-CN-GRI-XL",  30),
    ("2026-04-22", "Empresa TechCorp",           "TAZ-CER-BCA",   200),
    ("2026-04-22", "Empresa TechCorp",           "LLA-ACR-RED",   300),
    ("2026-04-25", "Casamiento Pérez-García",    "TAZ-CER-BCA",    80),
    ("2026-04-28", "Movistar - Día del Padre",   "REM-DF-NEG-M",   45),
    ("2026-04-28", "Movistar - Día del Padre",   "REM-DF-NEG-L",   35),
    ("2026-05-02", "Mercado Libre - Onboarding", "BUZ-CN-NEG-L",   50),
    ("2026-05-02", "Mercado Libre - Onboarding", "BUZ-CN-NEG-M",   40),
    ("2026-05-05", "Hospital Italiano",          "CHO-PIQ-BCO-M",  30),
    ("2026-05-08", "YPF - Estaciones",           "GOR-RES-AZU",    50),
    ("2026-05-08", "YPF - Estaciones",           "REM-DF-AZU-M",   25),
    ("2026-05-10", "Tienda Nube - Conferencia",  "TAZ-MAG",        60),
    ("2026-05-12", "Personal - Evento",          "CAM-CIE-NEG-L",  18),
    ("2026-05-14", "Globant - Welcome Kit",      "LLA-ACR-REC",   200),
    ("2026-05-14", "Globant - Welcome Kit",      "BUZ-CR-GRI-M",   30),
]

# ---------------------------------------------------------------------------
# Build workbook
# ---------------------------------------------------------------------------
wb = Workbook()

# --- HOJA STOCK ---
ws = wb.active
ws.title = "Stock"

HEADERS = ["SKU", "Producto", "Modelo", "Color", "Talle",
           "Stock Inicial", "Ingresos", "Pedidos", "Stock Actual", "Estado"]
ws.append(HEADERS)

header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
header_fill = PatternFill("solid", start_color="1F4E79")
header_align = Alignment(horizontal="center", vertical="center")
thin = Side(border_style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col in range(1, len(HEADERS) + 1):
    c = ws.cell(row=1, column=col)
    c.font = header_font
    c.fill = header_fill
    c.alignment = header_align
    c.border = border

for i, prod in enumerate(PRODUCTOS, start=2):
    sku, producto, modelo, color, talle, inicial, ingresos = prod
    ws.cell(row=i, column=1, value=sku)
    ws.cell(row=i, column=2, value=producto)
    ws.cell(row=i, column=3, value=modelo)
    ws.cell(row=i, column=4, value=color)
    ws.cell(row=i, column=5, value=talle)
    ws.cell(row=i, column=6, value=inicial)
    ws.cell(row=i, column=7, value=ingresos)
    ws.cell(row=i, column=8, value=f"=SUMIF(Pedidos!C:C,A{i},Pedidos!E:E)")
    ws.cell(row=i, column=9, value=f"=F{i}+G{i}-H{i}")
    ws.cell(row=i, column=10, value=f'=IF(I{i}>0,"OK","COMPRAR")')

for row in ws.iter_rows(min_row=2, max_row=len(PRODUCTOS)+1, max_col=len(HEADERS)):
    for c in row:
        c.border = border
        c.font = Font(name="Arial", size=10)
        if c.column >= 6:
            c.alignment = Alignment(horizontal="right")

widths = {"A": 16, "B": 12, "C": 22, "D": 16, "E": 8,
          "F": 13, "G": 11, "H": 10, "I": 13, "J": 11}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

verde = PatternFill("solid", start_color="C6EFCE")
rojo = PatternFill("solid", start_color="FFC7CE")
amarillo = PatternFill("solid", start_color="FFEB9C")

rango_stock = f"I2:I{len(PRODUCTOS)+1}"
ws.conditional_formatting.add(rango_stock, CellIsRule(operator="lessThan", formula=["0"], fill=rojo))
ws.conditional_formatting.add(rango_stock, CellIsRule(operator="equal", formula=["0"], fill=amarillo))
ws.conditional_formatting.add(rango_stock, CellIsRule(operator="greaterThan", formula=["0"], fill=verde))

rango_estado = f"J2:J{len(PRODUCTOS)+1}"
ws.conditional_formatting.add(rango_estado, FormulaRule(formula=['J2="COMPRAR"'], fill=rojo))
ws.conditional_formatting.add(rango_estado, FormulaRule(formula=['J2="OK"'], fill=verde))

ws.freeze_panes = "A2"

# --- HOJA PEDIDOS ---
ws2 = wb.create_sheet("Pedidos")
ws2.append(["Fecha", "Cliente", "SKU", "Producto", "Cantidad"])
for col in range(1, 6):
    c = ws2.cell(row=1, column=col)
    c.font = header_font
    c.fill = header_fill
    c.alignment = header_align
    c.border = border

sku_to_nombre = {p[0]: f"{p[1]} {p[2]} {p[3]} {p[4]}" for p in PRODUCTOS}

for i, pedido in enumerate(PEDIDOS, start=2):
    fecha, cliente, sku, cantidad = pedido
    ws2.cell(row=i, column=1, value=fecha)
    ws2.cell(row=i, column=2, value=cliente)
    ws2.cell(row=i, column=3, value=sku)
    ws2.cell(row=i, column=4, value=sku_to_nombre.get(sku, "?"))
    ws2.cell(row=i, column=5, value=cantidad)

for row in ws2.iter_rows(min_row=2, max_row=len(PEDIDOS)+1, max_col=5):
    for c in row:
        c.border = border
        c.font = Font(name="Arial", size=10)
        if c.column == 5:
            c.alignment = Alignment(horizontal="right")

ws2.column_dimensions["A"].width = 12
ws2.column_dimensions["B"].width = 32
ws2.column_dimensions["C"].width = 16
ws2.column_dimensions["D"].width = 38
ws2.column_dimensions["E"].width = 10
ws2.freeze_panes = "A2"

import os
output_path = os.path.join(os.path.dirname(__file__), "Stock_Merchandising_Demo.xlsx")
wb.save(output_path)
print(f"OK: {output_path}")
print(f"Productos en Stock: {len(PRODUCTOS)}")
print(f"Pedidos cargados: {len(PEDIDOS)}")
